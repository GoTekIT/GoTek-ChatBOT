from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
import re
root=Path('/Users/ngxuanphu/Documents/ChatGPT/GoTek ChatBOT - CTO')
src=(root/'reports/GOTEK_CHATBOT_DEV_HANDOFF.md').read_text()
d=Document(); sec=d.sections[0]; sec.page_height=Inches(11.69);sec.page_width=Inches(8.27)
sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.72);sec.right_margin=Inches(.72)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
 s=d.styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
 s.font.size=Pt({'Normal':10.5,'Title':23,'Heading 1':17,'Heading 2':12.5,'Heading 3':11}[name])
 s.paragraph_format.space_after=Pt(6);s.paragraph_format.line_spacing=1.08
for name in ['Heading 1','Heading 2','Heading 3']:
 d.styles[name].paragraph_format.space_before=Pt(10);d.styles[name].paragraph_format.keep_with_next=True
footer=sec.footer.paragraphs[0];footer.alignment=2
r=footer.add_run('GoTek ChatBOT  |  ');r.font.size=Pt(8)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
# Use native hyperlinks with wrapping opportunities in displayed labels.
def add_text(p,text):
 text=text.replace('`','')
 pattern=r'https://[^\s]+'
 pos=0
 for m in re.finditer(pattern,text):
  p.add_run(text[pos:m.start()]);u=m.group();rel=p.part.relate_to(u,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
  h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel);r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'2255AA');pr.append(c);r.append(pr);t=OxmlElement('w:t');t.text=u.replace('/','/\u200b').replace('%2F','%2F\u200b');r.append(t);h.append(r);p._p.append(h);pos=m.end()
 p.add_run(text[pos:])
def table(rows):
 t=d.add_table(rows=1, cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 n=len(rows[0]);widths={3:[1.25,2.7,2.75],4:[1.6,1.7,1.7,1.7]}.get(n,[6.7/n]*n)
 if n==4 and rows[0][0]=='Mã':widths=[.45,1.15,2.65,2.45]
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,vals in enumerate(rows):
  cells=t.rows[0].cells if i==0 else t.add_row().cells
  for j,(c,v) in enumerate(zip(cells,vals)):
   c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   p=c.paragraphs[0];p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.02
   r=p.add_run(v);r.font.size=Pt(9.3);r.bold=i==0
   tcpr=c._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E6EDF5' if i==0 else ('F6F8FA' if i%2==0 else 'FFFFFF'));tcpr.append(shade)
   mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');mar.append(e)
   tcpr.append(mar)
  trpr=t.rows[i]._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');trpr.append(ns)
  if i==0:
   rep=OxmlElement('w:tblHeader');trpr.append(rep)
 pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
 pr.append(b)
 d.add_paragraph().paragraph_format.space_after=Pt(0)
lines=src.splitlines();i=0;page=1
while i<len(lines):
 line=lines[i].strip();i+=1
 if not line:continue
 if line=='---PAGE---':page+=1;continue
 if line.startswith('|'):
  rows=[]
  while True:
   vals=[x.strip() for x in line.strip('|').split('|')]
   if not all(re.fullmatch(r'[-: ]+',v) for v in vals):rows.append(vals)
   if i>=len(lines) or not lines[i].strip().startswith('|'):break
   line=lines[i].strip();i+=1
  table(rows);continue
 if line.startswith('# '):p=d.add_paragraph(line[2:],'Title')
 elif line.startswith('## '):p=d.add_paragraph(line[3:],'Heading 1')
 elif line.startswith('### '):p=d.add_paragraph(line[4:],'Heading 2')
 elif line.startswith('#### '):p=d.add_paragraph(line[5:],'Heading 3')
 elif line.startswith('- '):p=d.add_paragraph(style='List Bullet');add_text(p,line[2:])
 else:p=d.add_paragraph();add_text(p,line)
 if page==12:
  p.paragraph_format.space_after=Pt(4)
  for r in p.runs:
   if p.style.name=='Normal':r.font.size=Pt(9)
# Add clickable contents for the thirteen main sections.
for n in range(1,14):
 targets=[p for p in d.paragraphs if p.style.name=='Heading 1' and p.text.startswith(str(n)+' ')]
 if not targets: continue
 target=targets[0]; start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(n));start.set(qn('w:name'),'section_'+str(n));target._p.insert(0,start)
 end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(n));target._p.append(end)
 toc=next((p for p in d.paragraphs if p.style.name=='Normal' and p.text.startswith(str(n)+'. ')),None)
 if toc is not None:
  label=toc.text;toc.clear();h=OxmlElement('w:hyperlink');h.set(qn('w:anchor'),'section_'+str(n));r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);toc._p.append(h)

for el in d.element.xpath('//w:pBdr'):
 el.getparent().remove(el)
for el in d.styles.element.xpath('//w:pBdr'):
 el.getparent().remove(el)
d.core_properties.title='Hồ sơ bàn giao phát triển nền tảng GoTek ChatBOT';d.core_properties.author='GoTek';d.core_properties.subject='Phân tích tài liệu và khảo sát HiChat'
out=root/'reports/GoTek_ChatBOT_Ban_giao_Dev_Toan_bo.docx';d.save(out);print(out)
